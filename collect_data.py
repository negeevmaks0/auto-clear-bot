import asyncio

from docx import Document



class CollectData:
    def __init__(self):
        self.suchente_text = 'Еженедельная уборка зала'

        self.column_mapping = {
            "Дата": "date_range",
            "Еженедельная уборка зала": "weekly_cleaning",
            "Уборка после встречи": "post_meeting_cleaning",
            "Ежемесячная уборка зала": "monthly_cleaning",
            "Гостеприимство": "hospitality"
        }

        self.cities_mapping = {
            'майнц - альтштадт': 'Altstadt',
            'майнц - кастeль':   'Kastel',
            'майнц - кастель':   'Kastel',
            'майнц - вайзенау':  'Weisenau',
            'майнц - гонзенхайм':'Gonsenheim',
            'майнц - лерхенберг':'Lerchenberg',
            'майнц - нойштадт':  'Neustadt',
            'майнц - центр':     'Zenter',
            'майнц - цeнтр':     'Zenter',
            'рюссельсхайм':      'Russelsheim'
        }


    def _clean_and_translate_city(self, text: str) -> str:
        if not text:
            return None
        
        clean_text = " ".join(text.strip().lower().split())
        
        return self.cities_mapping.get(clean_text, None)


    def _transform_date_to_db_key(self, date_range_str: str) -> str:
        months_map = {
            'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4, 'май': 5, 'июн': 6,
            'июл': 7, 'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12
        }

        try:
            clean_str = date_range_str.strip().lower()
            start_part = clean_str.split('-')[0].strip()
            month_name, day_str = start_part.split()
            
            day = int(day_str)
            month = months_map[month_name]
            
            import datetime
            today = datetime.date.today()
            current_year = today.year
            current_month = today.month
            
            if current_month >= 9 and month <= 4:
                calculated_year = current_year + 1

            elif current_month <= 4 and month >= 9:
                calculated_year = current_year - 1

            else:
                calculated_year = current_year

            short_year = str(calculated_year)[2:]
            
            return f"{short_year}.{month}.{day}"
            
        except Exception:
            return date_range_str


    def _parse_docx_with_merged_cells(self, download_path):
        doc = Document(download_path)
        
        target_table = None
        for table in doc.tables:
            if any("Дата" in cell.text for row in table.rows for cell in row.cells):
                target_table = table
                break
                
        if not target_table:
            return []

        headers = [cell.text.strip() for cell in target_table.rows[0].cells]
        db_headers = [self.column_mapping.get(h, h.lower().replace(" ", "_")) for h in headers]

        v_merged_values = {header: None for header in db_headers}
        parsed_data = []

        for row in target_table.rows[1:]:
            row_dict = {}
            
            for index, cell in enumerate(row.cells):
                column_name = db_headers[index]
                cell_value = cell.text.strip()
                
                tcPr = cell._tc.get_or_add_tcPr()
                vMerge = tcPr.find('{http://openxmlformats.org}vMerge')
                
                if vMerge is not None:
                    val = vMerge.get('{http://openxmlformats.org}val')

                    if val == "restart":
                        v_merged_values[column_name] = cell_value
                        raw_val = cell_value if cell_value != "" else None

                    else:
                        raw_val = v_merged_values[column_name] if v_merged_values[column_name] != "" else None

                else:
                    v_merged_values[column_name] = None
                    raw_val = cell_value if cell_value != "" else None

                if column_name == "date_range":
                    row_dict[column_name] = raw_val

                else:
                    row_dict[column_name] = self._clean_and_translate_city(raw_val)

            if row_dict.get("date_range"):
                row_dict["date_range"] = self._transform_date_to_db_key(row_dict["date_range"])
                parsed_data.append(row_dict)

        cleaned_data = []

        for item in parsed_data:
            filtered_item = {k: v for k, v in item.items() if k in self.column_mapping.values()}
            cleaned_data.append(filtered_item)

        return cleaned_data


    async def main(self, download_path):
        return await asyncio.to_thread(self._parse_docx_with_merged_cells, download_path)


    async def check_file(self, download_path):
        doc = Document(download_path)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = paragraph.text.strip()
                        if self.suchente_text in text:
                            return True

        return False

